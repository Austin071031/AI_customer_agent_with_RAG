"""
Comprehensive unit tests for Excel file upload related functions from Task 4.

This test suite covers:
- Excel data models validation
- SQLite database service operations
- Knowledge base manager Excel file processing
- File type detection and routing
- Error handling and edge cases
"""

import os
import tempfile
import pytest
import json
import sqlite3
from datetime import datetime
from pathlib import Path
from unittest.mock import Mock, patch, MagicMock

# Import the modules to test
from src.models.excel_models import (
    ExcelDocument, 
    ExcelSheetData, 
    ExcelFileUpload, 
    ExcelSearchQuery
)
from src.services.sqlite_database_service import SQLiteDatabaseService, SQLiteDatabaseError
from src.services.knowledge_base import (
    EnhancedKnowledgeBaseManager, 
    KnowledgeBaseError,
    DocumentProcessor,
    EmbeddingService
)


class TestExcelModels:
    """Test cases for Excel data models."""
    
    def test_excel_document_creation(self):
        """Test ExcelDocument model creation with valid data."""
        data = {
            "file_name": "sales_data.xlsx",
            "file_size": 10240,
            "sheet_names": ["Sales", "Customers", "Products"],
            "metadata": {
                "source": "company_reports",
                "uploaded_by": "admin"
            }
        }
        
        doc = ExcelDocument(**data)
        
        assert doc.file_name == "sales_data.xlsx"
        assert doc.file_size == 10240
        assert doc.sheet_names == ["Sales", "Customers", "Products"]
        assert doc.storage_type == "sqlite"
        assert doc.id.startswith("excel_")
        assert isinstance(doc.upload_time, datetime)
    
    def test_excel_document_invalid_file_extension(self):
        """Test ExcelDocument validation rejects invalid file extensions."""
        data = {
            "file_name": "sales_data.txt",
            "file_size": 10240,
            "sheet_names": ["Sheet1"]
        }
        
        with pytest.raises(ValueError, match="File name must have a valid Excel extension"):
            ExcelDocument(**data)
    
    def test_excel_document_empty_sheet_names(self):
        """Test ExcelDocument validation rejects empty sheet names."""
        data = {
            "file_name": "sales_data.xlsx",
            "file_size": 10240,
            "sheet_names": []
        }
        
        with pytest.raises(ValueError, match="Excel file must contain at least one sheet"):
            ExcelDocument(**data)
    
    def test_excel_sheet_data_creation(self):
        """Test ExcelSheetData model creation with valid data."""
        data = {
            "file_id": "excel_123e4567-e89b-12d3-a456-426614174000",
            "sheet_name": "Sales",
            "headers": ["Date", "Product", "Quantity", "Revenue"],
            "row_count": 150,
            "column_count": 4,
            "sample_data": [
                {"Date": "2024-01-01", "Product": "Widget A", "Quantity": 10, "Revenue": 1000.0},
                {"Date": "2024-01-02", "Product": "Widget B", "Quantity": 5, "Revenue": 750.0}
            ],
            "data_types": {
                "Date": "date",
                "Product": "string",
                "Quantity": "integer",
                "Revenue": "float"
            }
        }
        
        sheet_data = ExcelSheetData(**data)
        
        assert sheet_data.file_id == "excel_123e4567-e89b-12d3-a456-426614174000"
        assert sheet_data.sheet_name == "Sales"
        assert sheet_data.headers == ["Date", "Product", "Quantity", "Revenue"]
        assert sheet_data.row_count == 150
        assert sheet_data.column_count == 4
        assert len(sheet_data.sample_data) == 2
    
    def test_excel_sheet_data_invalid_file_id(self):
        """Test ExcelSheetData validation rejects invalid file_id format."""
        data = {
            "file_id": "invalid_id",
            "sheet_name": "Sales",
            "headers": ["Date", "Product"],
            "row_count": 10,
            "column_count": 2,
            "sample_data": [],
            "data_types": {}
        }
        
        with pytest.raises(ValueError, match='file_id must start with "excel_" prefix'):
            ExcelSheetData(**data)
    
    def test_excel_file_upload_creation(self):
        """Test ExcelFileUpload model creation with valid data."""
        data = {
            "file_name": "inventory_data.xlsx",
            "file_size": 20480,
            "description": "Current inventory levels",
            "tags": ["inventory", "stock"]
        }
        
        upload = ExcelFileUpload(**data)
        
        assert upload.file_name == "inventory_data.xlsx"
        assert upload.file_size == 20480
        assert upload.description == "Current inventory levels"
        assert upload.tags == ["inventory", "stock"]
    
    def test_excel_search_query_creation(self):
        """Test ExcelSearchQuery model creation with valid data."""
        data = {
            "query": "sales revenue",
            "file_id": "excel_123e4567-e89b-12d3-a456-426614174000",
            "sheet_name": "Sales",
            "column_name": "Product",
            "case_sensitive": False,
            "max_results": 50
        }
        
        search_query = ExcelSearchQuery(**data)
        
        assert search_query.query == "sales revenue"
        assert search_query.file_id == "excel_123e4567-e89b-12d3-a456-426614174000"
        assert search_query.sheet_name == "Sales"
        assert search_query.column_name == "Product"
        assert search_query.case_sensitive == False
        assert search_query.max_results == 50


class TestSQLiteDatabaseService:
    """Test cases for SQLite database service operations."""
    
    @pytest.fixture
    def temp_db(self):
        """Create a temporary database for testing."""
        with tempfile.NamedTemporaryFile(suffix='.db', delete=False) as f:
            db_path = f.name
        yield db_path
        # Cleanup
        if os.path.exists(db_path):
            os.unlink(db_path)
    
    @pytest.fixture
    def sqlite_service(self, temp_db):
        """Create SQLiteDatabaseService instance with temporary database."""
        return SQLiteDatabaseService(db_path=temp_db)
    
    @pytest.fixture
    def sample_excel_data(self):
        """Create sample Excel file data for testing."""
        return {
            "file_path": "test_data.xlsx",
            "file_name": "test_data.xlsx",
            "file_size": 1024,
            "sheet_names": ["Sheet1", "Sheet2"],
            "metadata": {
                "source": "test",
                "uploaded_by": "test_user"
            }
        }
    
    def test_database_initialization(self, sqlite_service):
        """Test that database tables are properly initialized."""
        # Verify tables exist by checking if we can query them
        with sqlite3.connect(sqlite_service.db_path) as conn:
            cursor = conn.cursor()
            
            # Check excel_documents table
            cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='excel_documents'")
            assert cursor.fetchone() is not None
            
            # Check excel_sheet_data table
            cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='excel_sheet_data'")
            assert cursor.fetchone() is not None
    
    def test_store_excel_file_success(self, sqlite_service, sample_excel_data):
        """Test successful storage of Excel file metadata."""
        # Mock the _process_and_store_sheet_data to avoid file I/O
        with patch.object(sqlite_service, '_process_and_store_sheet_data') as mock_process:
            file_id = sqlite_service.store_excel_file(**sample_excel_data)
            
            assert file_id.startswith("excel_")
            mock_process.assert_called_once()
    
    def test_store_excel_file_database_error(self, sqlite_service, sample_excel_data):
        """Test error handling when database operation fails."""
        with patch('sqlite3.connect') as mock_connect:
            mock_conn = Mock()
            mock_cursor = Mock()
            mock_conn.cursor.return_value = mock_cursor
            mock_cursor.execute.side_effect = sqlite3.Error("Database error")
            mock_connect.return_value.__enter__.return_value = mock_conn
            
            with pytest.raises(SQLiteDatabaseError, match="Excel file storage failed"):
                sqlite_service.store_excel_file(**sample_excel_data)
    
    def test_get_excel_file(self, sqlite_service, sample_excel_data):
        """Test retrieval of Excel file by ID."""
        # First store a file
        with patch.object(sqlite_service, '_process_and_store_sheet_data'):
            file_id = sqlite_service.store_excel_file(**sample_excel_data)
        
        # Then retrieve it
        retrieved_doc = sqlite_service.get_excel_file(file_id)
        
        assert retrieved_doc is not None
        assert retrieved_doc.id == file_id
        assert retrieved_doc.file_name == sample_excel_data["file_name"]
        assert retrieved_doc.file_size == sample_excel_data["file_size"]
    
    def test_get_nonexistent_excel_file(self, sqlite_service):
        """Test retrieval of non-existent Excel file returns None."""
        result = sqlite_service.get_excel_file("nonexistent_id")
        assert result is None
    
    def test_list_excel_files(self, sqlite_service, sample_excel_data):
        """Test listing all Excel files in database."""
        # Store multiple files
        with patch.object(sqlite_service, '_process_and_store_sheet_data'):
            file_ids = []
            for i in range(3):
                data = sample_excel_data.copy()
                data["file_name"] = f"file_{i}.xlsx"
                file_id = sqlite_service.store_excel_file(**data)
                file_ids.append(file_id)
        
        # List all files
        files = sqlite_service.list_excel_files()
        
        assert len(files) == 3
        assert all(isinstance(f, ExcelDocument) for f in files)
        assert all(f.id in file_ids for f in files)
    
    def test_delete_excel_file(self, sqlite_service, sample_excel_data):
        """Test deletion of Excel file from database."""
        # Store a file
        with patch.object(sqlite_service, '_process_and_store_sheet_data'):
            file_id = sqlite_service.store_excel_file(**sample_excel_data)
        
        # Delete it
        result = sqlite_service.delete_excel_file(file_id)
        
        assert result is True
        
        # Verify it's gone
        retrieved_doc = sqlite_service.get_excel_file(file_id)
        assert retrieved_doc is None
    
    def test_delete_nonexistent_excel_file(self, sqlite_service):
        """Test deletion of non-existent Excel file returns False."""
        result = sqlite_service.delete_excel_file("nonexistent_id")
        assert result is False
    
    def test_get_database_info(self, sqlite_service, sample_excel_data):
        """Test retrieval of database information."""
        # Store a file
        with patch.object(sqlite_service, '_process_and_store_sheet_data'):
            sqlite_service.store_excel_file(**sample_excel_data)
        
        info = sqlite_service.get_database_info()
        
        assert "document_count" in info
        assert "sheet_count" in info
        assert "total_size_bytes" in info
        assert info["document_count"] == 1
    
    def test_health_check(self, sqlite_service):
        """Test database health check."""
        result = sqlite_service.health_check()
        assert result is True
    
    def test_infer_data_type(self, sqlite_service):
        """Test data type inference from cell values."""
        test_cases = [
            (123, "integer"),
            (123.45, "float"),
            ("hello", "string"),
            ("2024-01-01", "date"),
            (True, "boolean"),
            (None, "null"),
            (datetime.now(), "datetime")
        ]
        
        for value, expected_type in test_cases:
            inferred_type = sqlite_service._infer_data_type(value)
            assert inferred_type == expected_type


class TestKnowledgeBaseExcelFunctions:
    """Test cases for knowledge base Excel file processing functions."""
    
    @pytest.fixture
    def temp_kb_dir(self):
        """Create temporary directory for knowledge base testing."""
        with tempfile.TemporaryDirectory() as temp_dir:
            yield temp_dir
    
    @pytest.fixture
    def kb_manager(self, temp_kb_dir):
        """Create KnowledgeBaseManager instance with temporary directory."""
        with patch('chromadb.PersistentClient') as mock_client:
            mock_collection = Mock()
            mock_client.return_value.get_or_create_collection.return_value = mock_collection
            mock_collection.count.return_value = 0
            
            manager = EnhancedKnowledgeBaseManager(
                persist_directory=temp_kb_dir,
                sqlite_db_path=os.path.join(temp_kb_dir, "test_excel.db")
            )
            manager.vector_store = mock_collection
            return manager
    
    @pytest.fixture
    def sample_excel_file(self):
        """Create a temporary Excel file for testing."""
        import openpyxl
        
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            # Create a simple Excel file
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "TestSheet"
            
            # Add headers and sample data
            sheet['A1'] = "Name"
            sheet['B1'] = "Age"
            sheet['C1'] = "City"
            
            sheet['A2'] = "Alice"
            sheet['B2'] = 30
            sheet['C2'] = "New York"
            
            sheet['A3'] = "Bob"
            sheet['B3'] = 25
            sheet['C3'] = "London"
            
            workbook.save(f.name)
            file_path = f.name
        
        yield file_path
        # Cleanup
        if os.path.exists(file_path):
            os.unlink(file_path)
    
    def test_detect_file_type_excel(self, kb_manager, sample_excel_file):
        """Test file type detection for Excel files."""
        file_type = kb_manager._detect_file_type(sample_excel_file)
        assert file_type == "excel"
    
    def test_detect_file_type_document(self, kb_manager):
        """Test file type detection for non-Excel files."""
        # Create a temporary text file
        with tempfile.NamedTemporaryFile(suffix='.txt', delete=False) as f:
            f.write(b"Sample text content")
            file_path = f.name
        
        try:
            file_type = kb_manager._detect_file_type(file_path)
            assert file_type == "document"
        finally:
            if os.path.exists(file_path):
                os.unlink(file_path)
    
    def test_add_documents_empty_list(self, kb_manager):
        """Test adding empty list of documents raises error."""
        with pytest.raises(KnowledgeBaseError, match="No file paths provided"):
            kb_manager.add_documents([])
    
    def test_add_documents_mixed_file_types(self, kb_manager, sample_excel_file):
        """Test adding mixed file types routes correctly."""
        # Create a temporary text file
        with tempfile.NamedTemporaryFile(suffix='.txt', delete=False) as f:
            f.write(b"Sample text content")
            text_file_path = f.name
        
        try:
            # Mock the processing methods
            with patch.object(kb_manager, '_process_excel_file') as mock_excel_process, \
                 patch.object(kb_manager, '_process_document_file') as mock_doc_process:
                
                mock_excel_process.return_value = {
                    'file_id': 'excel_test_id',
                    'file_name': 'test.xlsx',
                    'status': 'success'
                }
                mock_doc_process.return_value = {
                    'document_id': 'doc_test_id',
                    'file_name': 'test.txt',
                    'status': 'success'
                }
                
                results = kb_manager.add_documents([sample_excel_file, text_file_path])
                
                # Verify both processing methods were called
                mock_excel_process.assert_called_once()
                mock_doc_process.assert_called_once()
                
                # Verify results structure
                assert 'excel_files' in results
                assert 'documents' in results
                assert 'errors' in results
                assert len(results['excel_files']) == 1
                assert len(results['documents']) == 1
        
        finally:
            if os.path.exists(text_file_path):
                os.unlink(text_file_path)
    
    def test_process_excel_file_success(self, kb_manager, sample_excel_file):
        """Test successful processing of Excel file."""
        with patch.object(kb_manager.sqlite_service, 'store_excel_file') as mock_store:
            mock_store.return_value = "excel_test_id"
            
            result = kb_manager._process_excel_file(sample_excel_file)
            
            assert result['file_id'] == "excel_test_id"
            assert result['status'] == 'success'
            assert result['storage_type'] == 'sqlite'
            mock_store.assert_called_once()
    
    def test_process_excel_file_error(self, kb_manager, sample_excel_file):
        """Test error handling in Excel file processing."""
        with patch.object(kb_manager.sqlite_service, 'store_excel_file') as mock_store:
            mock_store.side_effect = Exception("Storage error")
            
            with pytest.raises(KnowledgeBaseError, match="Excel file processing failed"):
                kb_manager._process_excel_file(sample_excel_file)
    
    def test_search_excel_data(self, kb_manager):
        """Test searching Excel data."""
        search_query = ExcelSearchQuery(
            query="test query",
            file_id="excel_test_id",
            sheet_name="TestSheet"
        )
        
        expected_results = [
            {
                'file_name': 'test.xlsx',
                'sheet_name': 'TestSheet',
                'column_name': 'Name',
                'cell_value': 'Alice',
                'file_id': 'excel_test_id'
            }
        ]
        
        with patch.object(kb_manager.sqlite_service, 'search_excel_data') as mock_search:
            mock_search.return_value = expected_results
            
            results = kb_manager.search_excel_data(search_query)
            
            assert results == expected_results
            mock_search.assert_called_once_with(
                query="test query",
                file_id="excel_test_id",
                sheet_name="TestSheet",
                column_name=None,
                case_sensitive=False,
                max_results=50
            )
    
    def test_list_excel_files(self, kb_manager):
        """Test listing Excel files."""
        expected_files = [
            ExcelDocument(
                file_name="test1.xlsx",
                file_size=1024,
                sheet_names=["Sheet1"]
            ),
            ExcelDocument(
                file_name="test2.xlsx",
                file_size=2048,
                sheet_names=["Sheet1", "Sheet2"]
            )
        ]
        
        with patch.object(kb_manager.sqlite_service, 'list_excel_files') as mock_list:
            mock_list.return_value = expected_files
            
            files = kb_manager.list_excel_files()
            
            assert files == expected_files
            mock_list.assert_called_once()
    
    def test_get_excel_file(self, kb_manager):
        """Test retrieving specific Excel file."""
        expected_file = ExcelDocument(
            file_name="test.xlsx",
            file_size=1024,
            sheet_names=["Sheet1"]
        )
        
        with patch.object(kb_manager.sqlite_service, 'get_excel_file') as mock_get:
            mock_get.return_value = expected_file
            
            file = kb_manager.get_excel_file("excel_test_id")
            
            assert file == expected_file
            mock_get.assert_called_once_with("excel_test_id")
    
    def test_delete_excel_file(self, kb_manager):
        """Test deleting Excel file."""
        with patch.object(kb_manager.sqlite_service, 'delete_excel_file') as mock_delete:
            mock_delete.return_value = True
            
            result = kb_manager.delete_excel_file("excel_test_id")
            
            assert result is True
            mock_delete.assert_called_once_with("excel_test_id")
    
    def test_get_sheet_data(self, kb_manager):
        """Test retrieving sheet data."""
        expected_sheets = [
            ExcelSheetData(
                file_id="excel_test_id",
                sheet_name="Sheet1",
                headers=["Name", "Age"],
                row_count=10,
                column_count=2
            )
        ]
        
        with patch.object(kb_manager.sqlite_service, 'get_sheet_data') as mock_get_sheets:
            mock_get_sheets.return_value = expected_sheets
            
            sheets = kb_manager.get_sheet_data("excel_test_id", "Sheet1")
            
            assert sheets == expected_sheets
            mock_get_sheets.assert_called_once_with("excel_test_id", "Sheet1")
    
    def test_get_knowledge_base_info(self, kb_manager):
        """Test retrieving knowledge base information."""
        with patch.object(kb_manager.vector_store, 'count') as mock_count, \
             patch.object(kb_manager.sqlite_service, 'get_database_info') as mock_db_info, \
             patch.object(kb_manager.embedding_service, 'get_model_info') as mock_model_info:
            
            mock_count.return_value = 5
            mock_db_info.return_value = {
                'document_count': 3,
                'sheet_count': 5,
                'total_size_bytes': 10240
            }
            mock_model_info.return_value = {
                'model_name': 'test-model',
                'device': 'cpu'
            }
            
            info = kb_manager.get_knowledge_base_info()
            
            assert 'vector_store' in info
            assert 'sqlite_database' in info
            assert 'total_documents' in info
            assert info['total_documents'] == 8  # 5 + 3
    
    def test_health_check(self, kb_manager):
        """Test knowledge base health check."""
        with patch.object(kb_manager, 'get_knowledge_base_info') as mock_info, \
             patch.object(kb_manager.embedding_service, 'get_model_info') as mock_model_info, \
             patch.object(kb_manager.sqlite_service, 'health_check') as mock_sqlite_health:
            
            mock_info.return_value = {"test": "data"}
            mock_model_info.return_value = {"model": "info"}
            mock_sqlite_health.return_value = True
            
            result = kb_manager.health_check()
            
            assert result is True


class TestDocumentProcessor:
    """Test cases for document processor functionality."""
    
    @pytest.fixture
    def document_processor(self):
        """Create DocumentProcessor instance."""
        return DocumentProcessor()
    
    def test_supported_formats(self, document_processor):
        """Test that supported formats are correctly defined."""
        expected_formats = {'.pdf', '.txt', '.docx', '.doc', '.md'}
        assert document_processor.supported_formats == expected_formats
    
    def test_extract_text_from_nonexistent_file(self, document_processor):
        """Test error handling for non-existent files."""
        with pytest.raises(KnowledgeBaseError, match="File not found"):
            document_processor.extract_text_from_file("/nonexistent/file.txt")
    
    def test_extract_text_unsupported_format(self, document_processor):
        """Test error handling for unsupported file formats."""
        with tempfile.NamedTemporaryFile(suffix='.unsupported', delete=False) as f:
            f.write(b"test content")
            file_path = f.name
        
        try:
            with pytest.raises(KnowledgeBaseError, match="Unsupported file format"):
                document_processor.extract_text_from_file(file_path)
        finally:
            if os.path.exists(file_path):
                os.unlink(file_path)


class TestErrorHandling:
    """Test cases for error handling in Excel file upload functions."""
    
    def test_knowledge_base_error_creation(self):
        """Test KnowledgeBaseError creation with custom message."""
        error = KnowledgeBaseError("Test error message", "test_error")
        assert str(error) == "Test error message"
        assert error.message == "Test error message"
        assert error.error_type == "test_error"
    
    def test_sqlite_database_error_creation(self):
        """Test SQLiteDatabaseError creation with custom message."""
        error = SQLiteDatabaseError("Database error", "connection_error")
        assert str(error) == "Database error"
        assert error.message == "Database error"
        assert error.error_type == "connection_error"


if __name__ == "__main__":
    pass
